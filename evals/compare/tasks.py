"""The comparison task set: the SAME jobs, run through two harnesses on
the SAME model, judged by the SAME programmatic checks.

These are the agentic deliverables Nick cares about: spreadsheets and
CSVs, slide decks, single-file HTML applications with room for attention
to detail (the "os.html" and "solar system" genre of first-output tests),
plus code repair and multi-file work. Every task names an OUTPUT the
checks can open — a workbook openpyxl loads, a deck python-pptx loads, an
HTML file html.parser parses — and a list of FEATURES worth points. The
result is a score in [0, 1] per task, not a pass/fail, because "how
good" is the question a harness comparison asks.

Judging is deliberately mechanical (EVALS.md's rule: criteria you can
run). What it cannot see — whether the solar system is beautiful — the
report leaves to the person: every produced HTML file is linked so the
two harnesses' outputs can be opened side by side.
"""

from __future__ import annotations

import csv
import html.parser
import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

Check = Callable[[Path], tuple[bool, str]]


@dataclass
class Task:
    """One job both harnesses get, verbatim."""

    id: str
    category: str                       # office | web | code
    prompt: str
    # Files created in the workspace BEFORE the run (relative path → text).
    setup: dict[str, str] = field(default_factory=dict)
    # (name, check) pairs, each worth one point; the score is the fraction.
    checks: list[tuple[str, Check]] = field(default_factory=list)
    # Output files the report links for human review.
    artifacts: list[str] = field(default_factory=list)
    timeout_s: int = 600


# ------------------------------------------------------------------ helpers
def exists(rel: str) -> tuple[str, Check]:
    return (f"{rel} exists", lambda ws: ((ws / rel).exists(), rel))


def text_has(rel: str, *needles: str, any_of: bool = False) -> tuple[str, Check]:
    def check(ws: Path):
        p = ws / rel
        if not p.exists():
            return False, f"{rel} missing"
        text = p.read_text(encoding="utf-8", errors="replace").lower()
        hits = [n for n in needles if n.lower() in text]
        ok = bool(hits) if any_of else len(hits) == len(needles)
        return ok, f"found {hits} of {list(needles)}"
    return (f"{rel} has {needles}", check)


def regex_in(rel: str, pattern: str, label: str) -> tuple[str, Check]:
    def check(ws: Path):
        p = ws / rel
        if not p.exists():
            return False, f"{rel} missing"
        return bool(re.search(pattern, p.read_text(encoding="utf-8", errors="replace"),
                              re.IGNORECASE | re.DOTALL)), pattern
    return (label, check)


def min_size(rel: str, kb: int) -> tuple[str, Check]:
    return (f"{rel} >= {kb} KB",
            lambda ws: ((ws / rel).exists() and (ws / rel).stat().st_size >= kb * 1024,
                        f"{(ws / rel).stat().st_size if (ws / rel).exists() else 0} bytes"))


def html_parses(rel: str) -> tuple[str, Check]:
    """The file is well-formed enough for html.parser and has a <script>."""
    def check(ws: Path):
        p = ws / rel
        if not p.exists():
            return False, f"{rel} missing"
        text = p.read_text(encoding="utf-8", errors="replace")
        tags: list[str] = []

        class Counter(html.parser.HTMLParser):
            def handle_starttag(self, tag, attrs):
                tags.append(tag)
        try:
            Counter().feed(text)
        except Exception as error:
            return False, f"parse error: {error}"
        ok = "html" in tags and "script" in tags and "body" in tags
        return ok, f"{len(tags)} tags; script={'script' in tags}"
    return (f"{rel} is a complete HTML document with script", check)


def no_external_deps(rel: str) -> tuple[str, Check]:
    """Self-contained: no http(s) URLs in src/href attributes (CDN scripts
    are the usual leak; data: and # are fine)."""
    def check(ws: Path):
        p = ws / rel
        if not p.exists():
            return False, f"{rel} missing"
        text = p.read_text(encoding="utf-8", errors="replace")
        leaks = re.findall(r'(?:src|href)\s*=\s*["\']https?://[^"\']+', text, re.I)
        return not leaks, f"{len(leaks)} external references"
    return (f"{rel} is self-contained", check)


def xlsx_check(rel: str, fn: Callable, label: str) -> tuple[str, Check]:
    def check(ws: Path):
        p = ws / rel
        if not p.exists():
            return False, f"{rel} missing"
        try:
            import openpyxl
            wb = openpyxl.load_workbook(p)          # formulas kept as formulas
            return fn(wb)
        except Exception as error:
            return False, f"openpyxl: {error}"
    return (label, check)


def pptx_check(rel: str, fn: Callable, label: str) -> tuple[str, Check]:
    def check(ws: Path):
        p = ws / rel
        if not p.exists():
            return False, f"{rel} missing"
        try:
            from pptx import Presentation
            return fn(Presentation(str(p)))
        except Exception as error:
            return False, f"python-pptx: {error}"
    return (label, check)


def csv_check(rel: str, fn: Callable, label: str) -> tuple[str, Check]:
    def check(ws: Path):
        p = ws / rel
        if not p.exists():
            return False, f"{rel} missing"
        rows = list(csv.reader(io.StringIO(p.read_text(encoding="utf-8", errors="replace"))))
        return fn(rows)
    return (label, check)


# ------------------------------------------------------------------ the set
SALES_CSV = "\n".join(["region,product,units,unit_price",
                       "north,widget,120,2.50", "north,gadget,40,10.00",
                       "south,widget,300,2.50", "south,gizmo,15,45.00",
                       "east,gadget,75,10.00", "west,gizmo,9,45.00"]) + "\n"

TASKS: list[Task] = [
    # ---- office ---------------------------------------------------------
    Task(
        id="csv-revenue-column", category="office",
        prompt=("In the workspace there is sales.csv with columns region, product, "
                "units, unit_price. Add a revenue column (units × unit_price) to every "
                "row, then append a final row with region TOTAL, empty product and units, "
                "empty unit_price, and the sum of all revenue in the revenue column. "
                "Write the result to sales_with_revenue.csv. Verify it by reading the "
                "file back before you finish."),
        setup={"sales.csv": SALES_CSV},
        checks=[
            exists("sales_with_revenue.csv"),
            csv_check("sales_with_revenue.csv", lambda rows: (
                bool(rows) and "revenue" in [c.strip().lower() for c in rows[0]],
                f"header={rows[0] if rows else None}"), "has a revenue header"),
            csv_check("sales_with_revenue.csv", lambda rows: (
                any(r and r[0].strip().upper() == "TOTAL" for r in rows), "TOTAL row"),
                "has a TOTAL row"),
            csv_check("sales_with_revenue.csv", lambda rows: (
                any(r and r[0].strip().upper() == "TOTAL"
                    and abs(float((r[-1] or "0").replace(",", "")) - 3280.0) < 0.5 for r in rows),
                "total revenue should be 3280.00 (300+400+750+675+750+405)"), "total is 3280"),
        ],
        artifacts=["sales_with_revenue.csv"],
    ),
    Task(
        id="xlsx-budget-workbook", category="office",
        prompt=("Create budget.xlsx in the workspace with python (openpyxl is installed). "
                "Sheet 'Budget': header row Category, Jan, Feb, Mar, Total; five expense "
                "rows (Rent, Food, Transport, Utilities, Fun) with plausible numbers; the "
                "Total column must be Excel FORMULAS (=SUM(B2:D2) etc.), and a final row "
                "'Monthly total' with a SUM formula per month. Sheet 'Notes' with one "
                "sentence explaining the sheet. Bold the header row. Reopen the file with "
                "openpyxl to verify it loads before you finish."),
        checks=[
            exists("budget.xlsx"),
            xlsx_check("budget.xlsx", lambda wb: ("Budget" in wb.sheetnames and "Notes" in wb.sheetnames,
                                                  str(wb.sheetnames)), "two sheets Budget + Notes"),
            xlsx_check("budget.xlsx", lambda wb: (
                any(isinstance(c.value, str) and c.value.startswith("=SUM")
                    for row in wb["Budget"].iter_rows() for c in row) if "Budget" in wb.sheetnames else False,
                "formulas present"), "Total column uses SUM formulas"),
            xlsx_check("budget.xlsx", lambda wb: (
                "Budget" in wb.sheetnames and wb["Budget"].max_row >= 7 and wb["Budget"].max_column >= 5,
                f"{wb['Budget'].max_row if 'Budget' in wb.sheetnames else 0} rows"), "5 categories + totals"),
            xlsx_check("budget.xlsx", lambda wb: (
                "Budget" in wb.sheetnames and bool(wb["Budget"]["A1"].font and wb["Budget"]["A1"].font.bold),
                "A1 bold"), "header is bold"),
        ],
        artifacts=["budget.xlsx"],
    ),
    Task(
        id="pptx-five-slides", category="office",
        prompt=("Create seymour_pitch.pptx in the workspace with python (python-pptx is "
                "installed): a 5-slide deck about a local-first AI workspace called Seymour. "
                "Slide 1: a title slide with title and subtitle. Slides 2–4: a title and 3–5 "
                "bullet points each (the idea, how it works, why local). Slide 5: a closing "
                "slide with next steps. Add speaker notes to at least two slides. Reopen the "
                "file with python-pptx to verify it loads and has 5 slides before you finish."),
        checks=[
            exists("seymour_pitch.pptx"),
            pptx_check("seymour_pitch.pptx", lambda prs: (len(prs.slides) == 5, f"{len(prs.slides)} slides"),
                       "exactly 5 slides"),
            pptx_check("seymour_pitch.pptx", lambda prs: (
                all(any(sh.has_text_frame and sh.text_frame.text.strip() for sh in s.shapes) for s in prs.slides),
                "every slide has text"), "every slide has text"),
            pptx_check("seymour_pitch.pptx", lambda prs: (
                sum(1 for s in prs.slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip()) >= 2,
                "notes on >=2 slides"), "speaker notes on two slides"),
            pptx_check("seymour_pitch.pptx", lambda prs: (
                "seymour" in " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes
                                       if sh.has_text_frame).lower(), "mentions Seymour"), "on topic"),
        ],
        artifacts=["seymour_pitch.pptx"],
    ),
    # ---- web (single-file HTML, first output) -----------------------------
    Task(
        id="html-solar-system", category="web",
        prompt=("Create solar_system.html in the workspace: a single self-contained HTML "
                "file (vanilla HTML/CSS/JS, no external resources) that animates the solar "
                "system on a <canvas>: the sun, all eight planets orbiting at different "
                "speeds and distances with distinct colors and sizes, labels, Saturn's ring, "
                "a starfield background, a speed slider, and a pause/play button. Use "
                "requestAnimationFrame and make it look polished. Write the file; do not "
                "try to open a browser."),
        checks=[
            exists("solar_system.html"),
            html_parses("solar_system.html"),
            no_external_deps("solar_system.html"),
            regex_in("solar_system.html", r"<canvas", "uses a canvas"),
            regex_in("solar_system.html", r"requestAnimationFrame", "animates with requestAnimationFrame"),
            text_has("solar_system.html", "mercury", "venus", "earth", "mars", "jupiter",
                     "saturn", "uranus", "neptune"),
            regex_in("solar_system.html", r'type\s*=\s*["\']range["\']', "has a speed slider"),
            regex_in("solar_system.html", r"pause|play", "has pause/play"),
            min_size("solar_system.html", 6),
        ],
        artifacts=["solar_system.html"],
    ),
    Task(
        id="html-desktop-os", category="web",
        prompt=("Create os.html in the workspace: a single self-contained HTML file (vanilla "
                "HTML/CSS/JS, no external resources) that simulates a small desktop "
                "operating system in the browser: a wallpaper, a taskbar with a live clock, "
                "a start menu, desktop icons, and at least three apps that open in "
                "draggable, closable, minimizable windows (a notepad you can type in, a "
                "calculator that works, and an about window). Windows must come to the "
                "front when clicked. Make it look clean and modern. Write the file; do not "
                "try to open a browser."),
        checks=[
            exists("os.html"),
            html_parses("os.html"),
            no_external_deps("os.html"),
            regex_in("os.html", r"mousedown|pointerdown|dragstart", "draggable windows"),
            regex_in("os.html", r"setInterval|requestAnimationFrame", "live clock"),
            text_has("os.html", "calculator", "notepad", "about"),
            regex_in("os.html", r"z-?index", "window stacking"),
            regex_in("os.html", r"textarea|contenteditable", "typeable notepad"),
            min_size("os.html", 8),
        ],
        artifacts=["os.html"],
    ),
    # ---- code -------------------------------------------------------------
    Task(
        id="code-fix-and-test", category="code",
        prompt=("stats.py in the workspace has bugs: its own tests (test_stats.py) fail. "
                "Run the tests with `python -m pytest -q`, fix stats.py (not the tests) "
                "until they pass, and finish by stating what was wrong."),
        setup={
            "stats.py": ("def mean(xs):\n    return sum(xs) / (len(xs) - 1)\n\n\n"
                         "def median(xs):\n    s = sorted(xs)\n    n = len(s)\n"
                         "    return s[n // 2]\n\n\n"
                         "def variance(xs):\n    m = mean(xs)\n"
                         "    return sum((x - m) ** 2 for x in xs) / len(xs)\n"),
            "test_stats.py": ("from stats import mean, median, variance\n\n\n"
                              "def test_mean():\n    assert mean([2, 4, 6]) == 4\n\n\n"
                              "def test_median_even():\n    assert median([1, 3, 2, 4]) == 2.5\n\n\n"
                              "def test_median_odd():\n    assert median([3, 1, 2]) == 2\n\n\n"
                              "def test_variance():\n    assert abs(variance([2, 4, 4, 4, 5, 5, 7, 9]) - 4.0) < 1e-9\n"),
        },
        checks=[
            ("tests pass", lambda ws: _pytest_passes(ws)),
            text_has("test_stats.py", "def test_median_even"),      # tests untouched
        ],
        artifacts=["stats.py"],
    ),
    Task(
        id="code-multifile-refactor", category="code",
        prompt=("The workspace has app/config.py, app/db.py and app/main.py which all "
                "define their own copy of a function `read_env(name, default)`. Refactor: "
                "create app/env.py with the one canonical read_env (same behavior: return "
                "os.environ.get(name, default)), remove the copies, and make the three "
                "modules import it from app.env. Then run `python -m pytest -q` (tests "
                "exist in test_app.py) and make sure they pass before you finish."),
        setup={
            "app/__init__.py": "",
            "app/config.py": ("import os\n\n\ndef read_env(name, default=None):\n"
                              "    return os.environ.get(name, default)\n\n\n"
                              "DEBUG = read_env('APP_DEBUG', '0') == '1'\n"),
            "app/db.py": ("import os\n\n\ndef read_env(name, default=None):\n"
                          "    return os.environ.get(name, default)\n\n\n"
                          "DB_URL = read_env('APP_DB', 'sqlite://')\n"),
            "app/main.py": ("import os\n\n\ndef read_env(name, default=None):\n"
                            "    return os.environ.get(name, default)\n\n\n"
                            "def port():\n    return int(read_env('APP_PORT', '8000'))\n"),
            "test_app.py": ("import importlib\n\n\ndef test_single_definition():\n"
                            "    import app.env\n    from app import config, db, main\n"
                            "    assert config.read_env is app.env.read_env\n"
                            "    assert db.read_env is app.env.read_env\n"
                            "    assert main.read_env is app.env.read_env\n\n\n"
                            "def test_behavior(monkeypatch):\n"
                            "    monkeypatch.setenv('APP_PORT', '9000')\n"
                            "    from app import main\n    assert main.port() == 9000\n"),
        },
        checks=[
            exists("app/env.py"),
            ("tests pass", lambda ws: _pytest_passes(ws)),
            regex_in("app/config.py", r"from app\.env import|import app\.env", "config imports env"),
        ],
        artifacts=["app/env.py", "app/config.py"],
    ),
]


def _pytest_passes(ws: Path) -> tuple[bool, str]:
    """Run the workspace's tests with the project venv's python."""
    import subprocess
    import sys
    try:
        proc = subprocess.run([sys.executable, "-m", "pytest", "-q", "-p", "no:cacheprovider"],
                              cwd=ws, capture_output=True, text=True, timeout=120)
    except subprocess.TimeoutExpired:
        return False, "pytest timed out"
    tail = (proc.stdout or proc.stderr).strip().splitlines()[-1:] or ["(no output)"]
    return proc.returncode == 0, tail[0][:120]
