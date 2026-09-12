"""Build the calibration references and render them.

Three hand-scored artifacts the judge re-scores on every run (README:
a judge whose calibration moved is reporting noise). They are generated
here so they are reproducible and small enough to live in the repo as
code; `calibration.json` carries the scores fixed by hand after looking
at the renders, with the reasoning.

    .venv/bin/python evals/render/calibration/make.py
"""

from __future__ import annotations

import asyncio
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]
sys.path.insert(0, str(REPO))

from evals.render import render_artifact  # noqa: E402


def good_deck(path: Path) -> None:
    """A deck built the way the pptx skill teaches: one palette, a grid,
    short bullets, a chart from numbers, notes. Fixed score: 7."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt
    INK, BG, ACCENT = RGBColor(0x1F, 0x29, 0x37), RGBColor(0xFA, 0xFA, 0xF7), RGBColor(0x0E, 0x7C, 0x86)
    W, M = 13.333, 0.6
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(7.5)
    blank = prs.slide_layouts[6]

    def slide(title: str):
        s = prs.slides.add_slide(blank)
        s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
        box = s.shapes.add_textbox(Inches(M), Inches(0.4), Inches(W - 2 * M), Inches(0.9))
        p = box.text_frame.paragraphs[0]; p.text = title; p.font.size = Pt(32); p.font.bold = True; p.font.color.rgb = INK
        bar = s.shapes.add_shape(1, Inches(M), Inches(1.3), Inches(1.2), Inches(0.06))
        bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()
        return s

    def bullets(s, items, left=M, top=1.6, width=W - 2 * M, height=5.0):
        tf = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)).text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = "•  " + item; p.font.size = Pt(18); p.font.color.rgb = INK; p.space_after = Pt(10)

    s = prs.slides.add_slide(blank)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = INK
    t = s.shapes.add_textbox(Inches(M), Inches(2.6), Inches(W - 2 * M), Inches(1.4)).text_frame
    t.text = "Q3 sales review"; t.paragraphs[0].font.size = Pt(44); t.paragraphs[0].font.bold = True; t.paragraphs[0].font.color.rgb = BG
    s.notes_slide.notes_text_frame.text = "Welcome; three slides."
    s = slide("North carries the quarter")
    data = CategoryChartData(); data.categories = ["north", "south", "east", "west"]; data.add_series("Sales", [1420, 850, 580, 410])
    chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(M), Inches(1.6), Inches(8), Inches(5.2), data).chart
    chart.has_legend = False
    bullets(s, ["North is 44% of sales", "Four regions, 3,260 total"], left=9.0, top=1.8, width=3.7, height=4.5)
    s.notes_slide.notes_text_frame.text = "Lead with the concentration."
    s = slide("Three things to do next")
    bullets(s, ["Rebalance the pipeline toward east and west", "Repeat the north playbook", "Review gizmo pricing"])
    s.notes_slide.notes_text_frame.text = "Owners and dates follow."
    prs.save(path)


def bad_deck(path: Path) -> None:
    """The default template, a wall of text that overflows, an empty
    placeholder, no notes. Fixed score: 3."""
    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Sales"
    s.placeholders[1].text_frame.text = ("The quarter was characterised by a number of developments across the regions which we "
                                         "will now go through in considerable detail, region by region, product by product. ") * 6
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Next steps"
    prs.save(path)


def page(path: Path) -> None:
    """A small but complete single-file app: a clock, a working counter,
    responsive layout. Fixed score: 6."""
    path.write_text("""<!doctype html><html><head><meta name="viewport" content="width=device-width, initial-scale=1">
<style>body{margin:0;font:16px system-ui;background:#111827;color:#f9fafb;display:flex;flex-direction:column;min-height:100vh}
header{padding:16px 24px;background:#1f2937;display:flex;justify-content:space-between;align-items:center}
main{flex:1;display:grid;gap:16px;padding:24px;grid-template-columns:repeat(auto-fit,minmax(220px,1fr))}
.card{background:#1f2937;border-radius:12px;padding:20px}button{font:inherit;padding:8px 14px;border:0;border-radius:8px;background:#0e7c86;color:#fff}
</style></head><body><header><strong>Dashboard</strong><span id="clock">--:--:--</span></header>
<main><div class="card"><h2>Counter</h2><p id="count">0</p><button id="inc">Add one</button></div>
<div class="card"><h2>Notes</h2><textarea id="notes" rows="4" style="width:100%"></textarea></div>
<div class="card"><h2>About</h2><p>A calibration page for the judge.</p></div></main>
<script>let n=0;document.getElementById('inc').onclick=()=>{n++;document.getElementById('count').textContent=n};
setInterval(()=>{document.getElementById('clock').textContent=new Date().toLocaleTimeString()},1000);</script></body></html>""")


async def main() -> None:
    items = [
        ("deck-good", "review.pptx", good_deck, "pptx"),
        ("deck-bad", "sales.pptx", bad_deck, "pptx"),
        ("page-dashboard", "dashboard.html", page, "html"),
    ]
    for item_id, filename, builder, _kind in items:
        folder = HERE / item_id
        folder.mkdir(exist_ok=True)
        target = folder / filename
        builder(target)
        steps = ["click #inc", "type #notes hello"] if filename.endswith(".html") else None
        result = await render_artifact(target, folder / "render", steps)
        print(item_id, len(result["images"]), "images", result["notes"])


if __name__ == "__main__":
    asyncio.run(main())
