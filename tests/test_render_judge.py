"""The eval renderer and the judge packet: a deliverable becomes PNGs
deterministically, and the packet the judge sees is blind (random A/B,
neutral names, mapping kept aside) with the anchors verbatim."""

import json
from pathlib import Path

import pytest

from evals.judge import packet
from evals.render import render_artifact
from seymour.verify import office


def test_the_packet_is_blind_and_carries_the_anchors(tmp_path, monkeypatch):
    a = tmp_path / "seymour-out"; a.mkdir(); (a / "os-seymour.png").write_bytes(b"\x89PNG\r\n\x1a\nA")
    b = tmp_path / "dsh-out"; b.mkdir(); (b / "os-dsh.png").write_bytes(b"\x89PNG\r\n\x1a\nB")
    out = tmp_path / "judge" / "html-desktop-os"
    mapping = packet.build_packet(out, "Make os.html", "page",
                                  {"seymour": {"images": [str(a / "os-seymour.png")], "console": "no errors"},
                                   "dsh": {"images": [str(b / "os-dsh.png")], "console": ""}}, seed=3)
    assert set(mapping) == {"A", "B"} and set(mapping.values()) == {"seymour", "dsh"}
    prompt = (out / "prompt.md").read_text()
    for anchor in ("10 = professional enough to present to a client with no edits.",
                   "7 = correct and clean, visibly machine-made.",
                   "4 = the content is right, the presentation is not.",
                   "1 = broken or empty."):
        assert anchor in prompt
    assert "seymour" not in prompt.lower() and "dsh" not in prompt.lower()      # blind
    assert (out / "A" / "01.png").exists() and (out / "B" / "01.png").exists()   # neutral names
    assert "Make os.html" in prompt and "Would a person use this?" in prompt
    # A verdict folds back through the mapping.
    (out / "verdict.json").write_text(json.dumps({"artifacts": {"A": {"overall": 8, "highest_leverage_fix": "x", "criteria": []},
                                                               "B": {"overall": 5, "highest_leverage_fix": "y", "criteria": []}},
                                                 "winner": "A", "confidence": "high"}))
    monkeypatch.setattr(packet, "RESULTS", tmp_path)
    summary = packet.collect("judge".replace("judge", tmp_path.name) if False else "x") if False else None
    monkeypatch.setattr(packet, "RESULTS", tmp_path)
    (tmp_path / "judge-t").mkdir()
    (tmp_path / "judge-t" / "html-desktop-os").symlink_to(out)
    summary = packet.collect("t")
    row = summary["tasks"]["html-desktop-os"]
    assert row["winner"] == mapping["A"] and row[mapping["A"]]["overall"] == 8 and row[mapping["B"]]["overall"] == 5


@pytest.mark.skipif(not office.SOFFICE, reason="LibreOffice is not installed")
async def test_render_artifact_produces_slide_pngs_and_a_contact_sheet(tmp_path):
    from pptx import Presentation
    prs = Presentation()
    for title in ("One", "Two", "Three"):
        s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = title
    deck = tmp_path / "d.pptx"; prs.save(deck)
    out = await render_artifact(deck, tmp_path / "out")
    names = sorted(Path(p).name for p in out["images"])
    assert names == ["contact.png", "slide-01.png", "slide-02.png", "slide-03.png"] and out["notes"] == []
    missing = await render_artifact(tmp_path / "nope.pptx", tmp_path / "out2")
    assert missing["images"] == [] and "does not exist" in missing["notes"][0]
