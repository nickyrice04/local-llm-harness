"""The office verifiers and the page-interaction probe: a workbook is
judged by RECALCULATING it, a deck by rendering and by what cannot fit,
a document by loading and rendering — and a page by driving it. These
need LibreOffice and Playwright's Chromium; each test skips honestly
when its oracle is missing rather than passing by default."""

import asyncio
from pathlib import Path

import pytest

from seymour.config import settings
from seymour.tools import files, paths
from seymour.tools.verify import auto_check
from seymour.verify import office, web


@pytest.fixture(autouse=True)
def clean_workspace(tmp_path, monkeypatch):
    monkeypatch.setattr(settings.__class__, "workspace_dir",
                        property(lambda self: tmp_path / "ws"))
    (tmp_path / "ws").mkdir()
    files._seen.clear()
    yield


needs_soffice = pytest.mark.skipif(not office.SOFFICE, reason="LibreOffice (soffice) is not installed")


def _chromium_available() -> bool:
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as pw:
            pw.chromium.launch().close()
        return True
    except Exception:
        return False


needs_chromium = pytest.mark.skipif(not _chromium_available(), reason="Playwright's Chromium is not installed")


@needs_soffice
async def test_a_workbook_with_broken_formulas_fails_and_a_sound_one_passes_with_real_values():
    from openpyxl import Workbook
    ws_dir = paths.workspace()
    bad = Workbook(); sheet = bad.active; sheet.title = "Data"
    sheet["A1"], sheet["A2"], sheet["A3"] = 2, 3, "=SUM(A1:A2)"
    sheet["B1"] = "=A1/0"; sheet["C1"] = "=NOPE(1)"; sheet["D1"] = "=Missing!A1"
    bad.save(ws_dir / "bad.xlsx")
    verdict = await auto_check("bad.xlsx")
    assert verdict["kind"] == "xlsx" and verdict["verdict"] == "FIX NEEDED"
    assert "#DIV/0!" in verdict["report"] and "#NAME?" in verdict["report"] and "recalculated by LibreOffice" in verdict["report"]
    good = Workbook(); s = good.active; s.title = "Budget"
    s.append(["Category", "Jan", "Feb", "Total"]); s.append(["Rent", 100, 110, "=SUM(B2:C2)"]); s.append(["Food", 50, 60, "=SUM(B3:C3)"])
    s["A4"], s["B4"] = "Monthly total", "=SUM(B2:B3)"
    good.create_sheet("Notes")["A1"] = "hello"
    good.save(ws_dir / "good.xlsx")
    verdict = await auto_check("good.xlsx")
    assert verdict["verdict"] == "PASS", verdict["report"]
    assert "3 formula(s)" in verdict["report"] and "Budget!D2 = 210" in verdict["report"]   # a REAL recalculated value


@needs_soffice
async def test_a_deck_is_rendered_and_overflow_and_empty_placeholders_are_named():
    from pptx import Presentation
    ws_dir = paths.workspace()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Too much"
    slide.placeholders[1].text_frame.text = "words " * 400                      # cannot fit its box
    empty = prs.slides.add_slide(prs.slide_layouts[1])
    empty.shapes.title.text = "Empty body"                                       # body placeholder left blank
    prs.save(ws_dir / "deck.pptx")
    verdict = await auto_check("deck.pptx")
    assert verdict["kind"] == "pptx" and verdict["verdict"] == "FIX NEEDED"
    assert "overflows its box" in verdict["report"] and "slide 2: empty placeholder" in verdict["report"]
    assert "rendered 2 slide(s)" in verdict["report"]
    images = [Path(p) for p in verdict["images"]]
    assert len(images) == 2 and all(p.exists() and p.stat().st_size > 1000 for p in images)
    render_dir = paths.artifacts_dir() / "render" / "deck"
    assert (render_dir / "contact.png").exists() and (render_dir / "slide-01.png").exists()
    fine = Presentation()
    s = fine.slides.add_slide(fine.slide_layouts[1]); s.shapes.title.text = "Fine"; s.placeholders[1].text_frame.text = "one point"
    s.notes_slide.notes_text_frame.text = "say hello"
    fine.save(ws_dir / "fine.pptx")
    verdict = await auto_check("fine.pptx")
    assert verdict["verdict"] == "PASS" and "notes on 1" in verdict["report"], verdict["report"]


@needs_soffice
async def test_a_document_loads_and_renders():
    import docx
    document = docx.Document()
    document.add_heading("Report", level=1)
    document.add_paragraph("A paragraph of text. " * 20)
    document.add_table(rows=2, cols=2)
    document.save(paths.workspace() / "report.docx")
    verdict = await auto_check("report.docx")
    assert verdict["kind"] == "docx" and verdict["verdict"] == "PASS"
    assert "1 heading(s)" in verdict["report"] and "1 table(s)" in verdict["report"] and "rendered 1 page(s)" in verdict["report"]
    docx.Document().save(paths.workspace() / "empty.docx")
    assert (await auto_check("empty.docx"))["verdict"] == "FIX NEEDED"


PAGE = """<!doctype html><html><body>
<button id="go" onclick="document.getElementById('out').textContent='clicked'">Go</button>
<div id="out"></div>
<textarea id="pad"></textarea>
<div class="win" style="display:none"></div>
<div id="under" class="win" style="position:absolute;left:10px;top:120px;width:100px;height:60px;background:#999;z-index:1" onclick="this.style.zIndex=9"></div>
<div id="win" class="win" style="position:absolute;left:10px;top:120px;width:100px;height:60px;background:#ccc;z-index:5"><div class="bar">t</div></div>
<button id="dead">Dead</button>
<script>
const w = document.getElementById('win'); let drag = null;
w.addEventListener('mousedown', e => { drag = {x: e.clientX - w.offsetLeft, y: e.clientY - w.offsetTop}; });
window.addEventListener('mousemove', e => { if (drag) { w.style.left = (e.clientX - drag.x) + 'px'; w.style.top = (e.clientY - drag.y) + 'px'; } });
window.addEventListener('mouseup', () => { drag = null; });
</script></body></html>"""


@needs_chromium
async def test_interactions_report_whether_the_dom_changed_and_expectations_hold():
    page = paths.workspace() / "app.html"
    page.write_text(PAGE)
    steps = web.parse_steps("click #go; expect #out text clicked\ntype #pad hello world\nexpect #pad visible\n"
                            "drag .win .bar 80 40\nclick .win\nclick bottommost:.win\nexpect changed\nclick #dead\nexpect changed\nexpect #nope count 1")
    report = await web.run_interactions(page, steps)
    assert report is not None
    by_step = {s.step: s for s in report.steps}
    assert by_step["click #go"].ok and by_step["click #go"].changed is True
    assert by_step["expect #out text clicked"].ok
    assert by_step["type #pad hello world"].ok and by_step["expect #pad visible"].ok
    assert by_step["drag .win .bar 80 40"].ok and by_step["drag .win .bar 80 40"].changed is True   # a selector with a space; the window moved
    assert by_step["click .win"].ok                                            # hidden skipped; the topmost (z 5) is the target
    assert by_step["click bottommost:.win"].ok and by_step["click bottommost:.win"].changed is True   # the one underneath came forward
    assert by_step["click #dead"].changed is False                            # a button that does nothing
    assert by_step["expect changed"].ok is False                              # …and the assertion says so
    assert by_step["expect #nope count 1"].ok is False and "0 match" in by_step["expect #nope count 1"].note
    assert report.ok is False and report.console_errors == []
    # Through the tool: the interaction verdict rides on the load report.
    from seymour import tools
    out = await tools.execute("check_page", {"path": "app.html", "interact": "click #go\nexpect #out text clicked"})
    assert "[interaction] verdict: PASS" in out and "DOM changed" in out


@needs_chromium
async def test_the_screenshot_set_covers_viewports_scroll_and_interactions(tmp_path):
    page = paths.workspace() / "long.html"
    page.write_text(PAGE.replace("<div id=\"out\"></div>", "<div id=\"out\"></div>" + "<p>filler</p>" * 300))
    out = await web.screenshot_set(page, tmp_path / "shots", steps=["click #go", "type #pad hi"],
                                   viewports={"desktop": (1440, 900), "phone": (375, 812)})
    names = sorted(Path(p).name for p in out["images"])
    assert {"desktop-top.png", "desktop-middle.png", "desktop-bottom.png", "phone-top.png", "after-01.png", "after-02.png"} <= set(names)
    assert Path(out["console"]).exists() and out["interaction"] is not None and out["interaction"].ok


@needs_soffice
async def test_a_workbook_a_command_produced_is_verified_and_verify_file_works_on_demand(monkeypatch):
    """The real way deliverables are made — a script — reaches the verifier:
    run_command's result carries the workbook's recalculated verdict, and
    verify_file answers on demand without soffice in the sandbox."""
    from seymour import run_executor, tools
    from seymour.db import init_db
    init_db()
    import uuid
    log = run_executor.RunLog(str(uuid.uuid4()), run_executor.CHAT_POLICY)
    catalog = run_executor.catalog_for(run_executor.CHAT_POLICY)
    script = ("from openpyxl import Workbook\nwb=Workbook(); ws=wb.active; ws['A1']=1; ws['A2']='=A1/0'; wb.save('made.xlsx')")
    result = await run_executor._execute(log, catalog, "run_command", {"command": f"python -c \"{script}\""})
    assert "exit code 0" in result and "[auto-check made.xlsx] verdict: FIX NEEDED" in result and "#DIV/0!" in result
    assert log.last_check and log.last_check["path"] == "made.xlsx" and log.last_check["verdict"] == "FIX NEEDED"
    pages: dict = {}
    run_executor._track_page(pages, "run_command", {}, log)
    assert "made.xlsx" in run_executor._failing_pages(pages)               # the repair guard follows it
    on_demand = await tools.execute("verify_file", {"path": "made.xlsx"})
    assert "verdict: FIX NEEDED" in on_demand and "#DIV/0!" in on_demand
    assert "No verifier" in await tools.execute("verify_file", {"path": "orders.csv"}) or (await tools.execute("verify_file", {"path": "orders.csv"})).startswith("Error")
