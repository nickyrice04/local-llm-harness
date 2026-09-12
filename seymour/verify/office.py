"""Office verifiers: recalculate workbooks, render decks and documents.

LibreOffice (`soffice`, installed at /opt/homebrew/bin on this machine)
is the oracle: `--headless --convert-to xlsx` forces a full recalc of
every formula, and `--convert-to pdf` renders slides and pages the way
a person would see them. PyMuPDF rasterizes the PDF (pdftoppm is not
installed). Each soffice run gets its own profile directory, because a
shared profile serializes runs behind a lock file and a stale lock from
a killed run would wedge every later conversion.
"""

import asyncio
import math
import os
import shutil
import tempfile
import uuid
from pathlib import Path

from seymour.tools import paths

SOFFICE = shutil.which("soffice") or ("/opt/homebrew/bin/soffice" if os.path.exists("/opt/homebrew/bin/soffice") else "")
CONVERT_TIMEOUT_S = 180
RENDER_DPI = 150
# Excel's error literals, as openpyxl reads them back after a recalc.
ERROR_VALUES = ("#REF!", "#NAME?", "#DIV/0!", "#VALUE!", "#N/A", "#NUM!", "#NULL!")
# A slide's text is estimated to overflow when the lines it needs at its
# font size exceed the box's height by this factor (the estimate is
# rough — average glyph width — so a margin keeps false alarms down).
OVERFLOW_TOLERANCE = 1.15


async def soffice_convert(src: Path, fmt: str, outdir: Path, timeout: float = CONVERT_TIMEOUT_S) -> Path | None:
    """`soffice --headless --convert-to <fmt>` into outdir; the output
    path, or None when soffice is missing, fails or times out."""
    if not SOFFICE:
        return None
    outdir.mkdir(parents=True, exist_ok=True)
    profile = Path(tempfile.mkdtemp(prefix="seymour-soffice-"))
    argv = [SOFFICE, f"-env:UserInstallation=file://{profile}", "--headless", "--norestore",
            "--convert-to", fmt, "--outdir", str(outdir), str(src)]
    try:
        proc = await asyncio.create_subprocess_exec(*argv, stdout=asyncio.subprocess.PIPE,
                                                    stderr=asyncio.subprocess.STDOUT)
        try:
            await asyncio.wait_for(proc.wait(), timeout=timeout)
        except asyncio.TimeoutError:
            proc.kill()
            return None
    except OSError:
        return None
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    out = outdir / (src.stem + "." + fmt.split(":")[0])
    return out if out.exists() else None


def render_pdf(pdf: Path, outdir: Path, stem: str = "page", dpi: int = RENDER_DPI) -> list[Path]:
    """One PNG per page at `dpi`, named <stem>-01.png …"""
    import pymupdf
    outdir.mkdir(parents=True, exist_ok=True)
    images: list[Path] = []
    with pymupdf.open(pdf) as doc:
        for index, page in enumerate(doc, 1):
            target = outdir / f"{stem}-{index:02d}.png"
            page.get_pixmap(dpi=dpi).save(target)
            images.append(target)
    return images


def contact_sheet(images: list[Path], target: Path, columns: int = 3, cell_width: int = 480) -> Path | None:
    """Every page as a thumbnail on one image, so consistency across a
    deck can be judged at a glance. Built with PyMuPDF alone (a new PDF
    page with the thumbnails placed on a grid, rasterized once)."""
    import pymupdf
    if not images:
        return None
    first = pymupdf.Pixmap(str(images[0]))
    ratio = first.height / max(first.width, 1)
    cell_height = int(cell_width * ratio)
    rows = math.ceil(len(images) / columns)
    gap = 16
    width = columns * cell_width + (columns + 1) * gap
    height = rows * cell_height + (rows + 1) * gap
    doc = pymupdf.open()
    page = doc.new_page(width=width, height=height)
    for index, image in enumerate(images):
        col, row = index % columns, index // columns
        x = gap + col * (cell_width + gap)
        y = gap + row * (cell_height + gap)
        page.insert_image(pymupdf.Rect(x, y, x + cell_width, y + cell_height), filename=str(image))
        page.insert_text((x + 4, y + 14), str(index + 1), fontsize=11, color=(0.8, 0.1, 0.1))
    page.get_pixmap(dpi=96).save(target)
    doc.close()
    return target


def _render_dir(src: Path) -> Path:
    """Where a file's renders live: .seymour/artifacts/render/<name>/,
    cleared per render so stale pages never mix with fresh ones."""
    target = paths.artifacts_dir() / "render" / src.stem
    shutil.rmtree(target, ignore_errors=True)
    target.mkdir(parents=True, exist_ok=True)
    return target


# --------------------------------------------------------------------------- #
#  Workbooks                                                                   #
# --------------------------------------------------------------------------- #

async def recalc_xlsx(path: Path) -> dict:
    """Recalculate through soffice; read the values back; fail on error
    literals and on formulas that produced nothing."""
    import openpyxl
    rel = paths.display(path)
    try:
        formulas_wb = openpyxl.load_workbook(path)           # formulas as text
    except Exception as error:
        return {"kind": "xlsx", "verdict": "FIX NEEDED", "report": f"[auto-check {rel}] openpyxl cannot open it: {error}", "images": []}
    formula_cells: dict[str, list[str]] = {}
    for ws in formulas_wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_cells.setdefault(ws.title, []).append(cell.coordinate)
    n_formulas = sum(len(v) for v in formula_cells.values())
    if not SOFFICE:
        return {"kind": "xlsx", "verdict": "not measured",
                "report": f"[auto-check {rel}] opens with openpyxl ({len(formulas_wb.sheetnames)} sheets, {n_formulas} formulas) — "
                          "soffice is not installed, so the formulas could not be recalculated.", "images": []}
    with tempfile.TemporaryDirectory(prefix="seymour-recalc-") as tmp:
        out = await soffice_convert(path, "xlsx", Path(tmp))
        if out is None:
            return {"kind": "xlsx", "verdict": "not measured",
                    "report": f"[auto-check {rel}] soffice could not recalculate the workbook (conversion failed or timed out).", "images": []}
        try:
            values_wb = openpyxl.load_workbook(out, data_only=True)
        except Exception as error:
            return {"kind": "xlsx", "verdict": "FIX NEEDED", "report": f"[auto-check {rel}] the recalculated workbook does not load: {error}", "images": []}
        errors: list[str] = []
        empty: list[str] = []
        samples: list[str] = []
        for sheet, coords in formula_cells.items():
            if sheet not in values_wb.sheetnames:
                continue
            ws = values_wb[sheet]
            for coord in coords:
                value = ws[coord].value
                if isinstance(value, str) and value.strip() in ERROR_VALUES:
                    errors.append(f"{sheet}!{coord} = {value}")
                elif value is None:
                    empty.append(f"{sheet}!{coord}")
                elif len(samples) < 6:
                    samples.append(f"{sheet}!{coord} = {value!r} ({formulas_wb[sheet][coord].value})")
    facts = (f"{len(formulas_wb.sheetnames)} sheet(s) {formulas_wb.sheetnames}; {n_formulas} formula(s); "
             f"recalculated by LibreOffice")
    if errors or empty:
        lines = [f"[auto-check {rel}] verdict: FIX NEEDED — {facts}"]
        if errors:
            lines.append(f"formula errors ({len(errors)}): " + "; ".join(errors[:12]) + (" …" if len(errors) > 12 else ""))
        if empty:
            lines.append(f"formulas that produce NOTHING ({len(empty)}): " + ", ".join(empty[:12]) + (" …" if len(empty) > 12 else ""))
        lines.append("Fix the referenced ranges/names, then save again (the check reruns).")
        return {"kind": "xlsx", "verdict": "FIX NEEDED", "report": "\n".join(lines), "images": []}
    report = f"[auto-check {rel}] verdict: PASS — {facts}"
    if samples:
        report += "\nfacts: " + "; ".join(samples)
    elif n_formulas == 0:
        report += "\nfacts: no formulas in this workbook (values only) — if totals were asked for, they should be formulas."
    return {"kind": "xlsx", "verdict": "PASS", "report": report, "images": []}


# --------------------------------------------------------------------------- #
#  Decks                                                                       #
# --------------------------------------------------------------------------- #

def _text_overflows(shape) -> tuple[bool, str]:
    """Estimate whether a text frame's content fits its box: lines needed
    (at ~0.5 em per glyph, the frame's width) × line height (1.2 × font
    size) against the shape's height. Rough on purpose; the render is
    what a judge looks at, this is the early warning."""
    from pptx.util import Emu
    if not shape.has_text_frame or not shape.width or not shape.height:
        return False, ""
    text_frame = shape.text_frame
    width_pt = Emu(shape.width).pt
    height_pt = Emu(shape.height).pt
    needed_pt = 0.0
    for paragraph in text_frame.paragraphs:
        size = None
        for run in paragraph.runs:
            if run.font.size:
                size = run.font.size.pt
                break
        if size is None and paragraph.font.size:
            size = paragraph.font.size.pt
        size = size or 18.0
        text = "".join(run.text for run in paragraph.runs) or paragraph.text or ""
        chars_per_line = max(int(width_pt / (size * 0.5)), 1)
        lines = max(math.ceil(len(text) / chars_per_line), 1) if text.strip() else 1
        needed_pt += lines * size * 1.2
    if needed_pt > height_pt * OVERFLOW_TOLERANCE and text_frame.text.strip():
        return True, f"needs ~{needed_pt:.0f}pt, box is {height_pt:.0f}pt"
    return False, ""


async def check_pptx(path: Path, render: bool = True) -> dict:
    """Mechanical checks with python-pptx, then a render to PNGs."""
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER
    rel = paths.display(path)
    try:
        prs = Presentation(str(path))
    except Exception as error:
        return {"kind": "pptx", "verdict": "FIX NEEDED", "report": f"[auto-check {rel}] python-pptx cannot open it: {error}", "images": []}
    problems: list[str] = []
    count = len(prs.slides)
    if count == 0:
        problems.append("the deck has no slides")
    notes = 0
    for number, slide in enumerate(prs.slides, 1):
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            notes += 1
        for shape in slide.shapes:
            if shape.is_placeholder and shape.has_text_frame and not shape.text_frame.text.strip():
                # An empty text-bearing placeholder renders as "Click to add
                # text". A content (OBJECT) placeholder still HAS a text
                # frame only while nothing was put in it — a picture or a
                # table replaces it — so an empty one is empty.
                kind = shape.placeholder_format.type
                if kind not in (PP_PLACEHOLDER.PICTURE, PP_PLACEHOLDER.CHART, PP_PLACEHOLDER.TABLE, PP_PLACEHOLDER.MEDIA_CLIP):
                    problems.append(f"slide {number}: empty placeholder '{shape.name}' (\"Click to add …\" would show) — fill or delete it")
            overflow, why = _text_overflows(shape)
            if overflow:
                problems.append(f"slide {number}: text in '{shape.name}' overflows its box ({why}) — fewer words, a smaller font, or a bigger box")
    images: list[Path] = []
    render_note = ""
    if render and count and SOFFICE:
        outdir = _render_dir(path)
        pdf = await soffice_convert(path, "pdf", outdir)
        if pdf is not None:
            try:
                images = render_pdf(pdf, outdir, stem="slide")
                sheet = contact_sheet(images, outdir / "contact.png")
                render_note = (f"rendered {len(images)} slide(s) to {paths.display(outdir)}/slide-NN.png"
                               + (f" and {paths.display(sheet)}" if sheet else ""))
                if len(images) != count:
                    problems.append(f"python-pptx counts {count} slides but the render has {len(images)} pages")
            except Exception as error:                  # a render failure is a fact, not a crash
                render_note = f"render failed: {error}"
        else:
            render_note = "render not measured (soffice conversion failed)"
    elif render and count:
        render_note = "render not measured (soffice is not installed)"
    facts = f"{count} slide(s); notes on {notes}; {render_note}".strip("; ")
    if problems:
        report = (f"[auto-check {rel}] verdict: FIX NEEDED — {facts}\n" + "\n".join(f"- {p}" for p in problems[:15])
                  + ("\n…" if len(problems) > 15 else ""))
        return {"kind": "pptx", "verdict": "FIX NEEDED", "report": report, "images": [str(i) for i in images]}
    return {"kind": "pptx", "verdict": "PASS", "report": f"[auto-check {rel}] verdict: PASS — {facts}\nfacts: {facts}",
            "images": [str(i) for i in images]}


# --------------------------------------------------------------------------- #
#  Documents                                                                   #
# --------------------------------------------------------------------------- #

async def check_docx(path: Path, render: bool = True) -> dict:
    """python-docx loads it; count structure; render pages."""
    import docx
    rel = paths.display(path)
    try:
        document = docx.Document(str(path))
    except Exception as error:
        return {"kind": "docx", "verdict": "FIX NEEDED", "report": f"[auto-check {rel}] python-docx cannot open it: {error}", "images": []}
    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    headings = [p for p in paragraphs if p.style is not None and p.style.name.lower().startswith("heading")]
    tables = len(document.tables)
    words = sum(len(p.text.split()) for p in paragraphs)
    problems: list[str] = []
    if not paragraphs and not tables:
        problems.append("the document is empty")
    images: list[Path] = []
    render_note = ""
    if render and SOFFICE and (paragraphs or tables):
        outdir = _render_dir(path)
        pdf = await soffice_convert(path, "pdf", outdir)
        if pdf is not None:
            try:
                images = render_pdf(pdf, outdir, stem="page")
                render_note = f"rendered {len(images)} page(s) to {paths.display(outdir)}/page-NN.png"
            except Exception as error:
                render_note = f"render failed: {error}"
        else:
            render_note = "render not measured (soffice conversion failed)"
    facts = f"{len(paragraphs)} paragraph(s), {len(headings)} heading(s), {tables} table(s), {words} words; {render_note}".strip("; ")
    if problems:
        return {"kind": "docx", "verdict": "FIX NEEDED", "report": f"[auto-check {rel}] verdict: FIX NEEDED — {facts}\n- " + "\n- ".join(problems),
                "images": [str(i) for i in images]}
    return {"kind": "docx", "verdict": "PASS", "report": f"[auto-check {rel}] verdict: PASS — {facts}\nfacts: {facts}", "images": [str(i) for i in images]}


async def render_xlsx(path: Path) -> list[Path]:
    """A workbook as pages (recalculated → PDF → PNG): a spreadsheet is a
    document too, and it can look bad. For the eval renderer."""
    if not SOFFICE:
        return []
    outdir = _render_dir(path)
    with tempfile.TemporaryDirectory(prefix="seymour-recalc-") as tmp:
        recalculated = await soffice_convert(path, "xlsx", Path(tmp))
        source = recalculated or path
        pdf = await soffice_convert(source, "pdf", outdir)
        if pdf is None:
            return []
        return render_pdf(pdf, outdir, stem="sheet")


__all__ = ["SOFFICE", "soffice_convert", "render_pdf", "contact_sheet", "recalc_xlsx", "check_pptx", "check_docx", "render_xlsx"]
uuid  # noqa: B018  (kept: callers may name renders by uuid)
